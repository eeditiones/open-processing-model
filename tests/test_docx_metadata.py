"""ODD ``behaviour="metadata"`` → Word core document properties."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from opm.resources import packaged_odd
from opm.runtime.docx_output_functions import DocxOutputFunctions

_JATS = (
    '<article>'
    '<front><journal-meta><journal-title-group><journal-title>A Journal</journal-title>'
    '</journal-title-group></journal-meta>'
    '<article-meta>'
    '<title-group><article-title>An Article</article-title></title-group>'
    '<contrib-group>'
    '<contrib><name><surname>Lovelace</surname><given-names>Ada</given-names></name></contrib>'
    '<contrib><name><surname>Babbage</surname><given-names>Charles</given-names></name></contrib>'
    '</contrib-group>'
    '<abstract><title>Abstract</title><p>What the article covers.</p></abstract>'
    '</article-meta></front>'
    '<body><sec id="s1"><title>One</title><p>Body text.</p></sec></body>'
    '</article>'
)


def _docx(tmp_path: Path, odd: str, xml: str) -> Document:
    from opm.odd_compiler import compile_odd
    from opm.transform import load_transform_module, run_transform

    mod_path = tmp_path / f'{odd}_docx.py'
    mod_path.write_text(
        compile_odd(str(packaged_odd(odd)), output_mode='docx'), encoding='utf-8'
    )
    out = run_transform(
        load_transform_module(mod_path),
        etree.fromstring(xml.encode()),
        xpath_extensions=['opm.runtime.common_xpath_functions'],
    )
    assert isinstance(out, bytes)
    path = tmp_path / f'{odd}.docx'
    path.write_bytes(out)
    return Document(str(path))


def test_jats_docx_sets_properties_and_keeps_the_front_matter(tmp_path: Path) -> None:
    """A journal article needs both.

    ``behaviour="metadata"`` emits nothing, so the ODD pairs it with the display
    model in a ``modelSequence``: dropping the front matter from the body would
    leave the docx opening on its first section with the title only in
    File → Properties.
    """
    doc = _docx(tmp_path, 'jats', _JATS)

    props = doc.core_properties
    assert props.title == 'An Article'
    assert props.author == 'Ada Lovelace, Charles Babbage'
    assert props.comments == 'What the article covers.'

    body = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert 'An Article' in body
    assert any('Ada Lovelace' in line for line in body)
    assert any('What the article covers.' in line for line in body)


def test_over_long_property_is_truncated_not_fatal() -> None:
    """OOXML caps core properties at 255 chars and python-docx raises rather than clamp.

    Abstracts routinely run longer, so the whole transform used to fail.
    """
    doc = Document()
    pmf = DocxOutputFunctions()
    config = {'parameters': {'metadata': {'abstract': ['x' * 400]}}}
    pmf._apply_core_properties(config, doc)

    assert len(doc.core_properties.comments) == 255
    assert doc.core_properties.comments.endswith('…')


def test_unmapped_metadata_key_is_ignored() -> None:
    doc = Document()
    pmf = DocxOutputFunctions()
    pmf._apply_core_properties({'parameters': {'metadata': {'nonesuch': ['v']}}}, doc)
    assert doc.core_properties.title == ''


@pytest.mark.parametrize('key', ['title', 'authors', 'abstract'])
def test_core_property_keys_cover_the_odd_metadata_keys(key: str) -> None:
    assert key in DocxOutputFunctions._CORE_PROPERTY_KEYS
